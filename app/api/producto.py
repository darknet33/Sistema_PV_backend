from fastapi import APIRouter, Depends, HTTPException, UploadFile, File
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session
from typing import List
from io import BytesIO
import os
from pathlib import Path
from openpyxl import Workbook, load_workbook
from jose import jwt
from PIL import Image
from app.database import get_db
from app.schemas.producto import ProductoCreate, ProductoUpdate, ProductoResponse
from app.crud.producto import get_productos_full, get_producto, get_producto_full, get_producto_by_codigo, create_producto, update_producto, delete_producto, delete_productos_batch as crud_delete_batch, delete_all_productos as crud_delete_all
from app.models.producto import Producto
from app.models.categoria import Categoria
from app.models.usuario import Usuario
from app.auth import SECRET_KEY, ALGORITHM, oauth2_scheme, get_current_user_full
from app.ws import broadcast_sync, broadcast_multiple_sync

router = APIRouter()

BASE_DIR = Path(__file__).resolve().parent.parent.parent
UPLOAD_DIR = BASE_DIR / "uploads" / "productos"
ALLOWED_EXTENSIONS = {".jpg", ".jpeg", ".png", ".gif", ".webp"}

@router.get("/", response_model=List[ProductoResponse])
def read_productos(skip: int = 0, limit: int = 10000, db: Session = Depends(get_db)):
    return get_productos_full(db, skip, limit)

@router.get("/export-xlsx")
def export_productos(db: Session = Depends(get_db)):
    productos = db.query(Producto).join(Categoria, isouter=True).all()
    
    wb = Workbook()
    ws = wb.active
    ws.title = "Productos"
    
    headers = ["ÏD", "Código", "Categoría", "Descripción", "Marca", "Procedencia", "Costo Bs.", "Utilidad Bs.", "Stock Inicial", "Stock Mínimo", "Stock Máximo", "Estado"]
    ws.append(headers)
    
    for p in productos:
        ws.append([
            p.id,
            p.codigo,
            p.categoria.nombre if p.categoria else "",
            p.descripcion,
            p.marca,
            p.procedencia,
            float(p.precio),
            float(p.utilidad),
            p.stock_inicial,
            p.stock_minimo,
            p.stock_maximo,
            "Activo" if p.activo else "Inactivo",
        ])
    
    output = BytesIO()
    wb.save(output)
    output.seek(0)
    
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": 'attachment; filename="productos.xlsx"'}
    )

@router.post("/import-xlsx")
async def import_productos(
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
    token: str = Depends(oauth2_scheme),
):
    user = None
    try:
        payload = jwt.decode(token, SECRET_KEY, algorithms=[ALGORITHM])
        username: str = payload.get("sub")
        if username:
            user = db.query(Usuario).filter(Usuario.username == username).first()
    except Exception:
        pass
    
    usuario_id = user.id if user else None
    
    if not file.filename.endswith((".xlsx", ".xls")):
        raise HTTPException(status_code=400, detail="Solo se permiten archivos Excel (.xlsx, .xls)")
    
    content = await file.read()
    wb = load_workbook(BytesIO(content))
    ws = wb.active
    
    resultados = {"creados": 0, "actualizados": 0, "errores": [], "procesados": 0}
    
    for row_num, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if not row or all(v is None for v in row):
            continue
        if len(row) < 12:
            resultados["errores"].append(f"Fila {row_num}: Solo {len(row)} columnas (se esperan 12)")
            continue
        
        try:
            _, codigo, categoria_nombre, descripcion, marca, procedencia, costo_bs, utilidad_bs, stock_inicial, stock_minimo, stock_maximo, estado = row
            
            if not codigo or not descripcion or not marca:
                resultados["errores"].append(f"Fila {row_num}: Faltan campos obligatorios (codigo, descripcion, marca)")
                continue
            
            resultados["procesados"] += 1
            
            categoria = None
            if categoria_nombre:
                cat_nombre = str(categoria_nombre).strip()
                categoria = db.query(Categoria).filter(Categoria.nombre == cat_nombre).first()
                if not categoria:
                    categoria = Categoria(nombre=cat_nombre)
                    db.add(categoria)
                    db.flush()
            
            producto_existente = db.query(Producto).filter(Producto.codigo == str(codigo).strip()).first()
            
            if producto_existente:
                producto_existente.categoria_id = categoria.id if categoria else producto_existente.categoria_id
                producto_existente.descripcion = str(descripcion)
                producto_existente.marca = str(marca)
                producto_existente.procedencia = str(procedencia) if procedencia else ""
                producto_existente.precio = float(costo_bs) if costo_bs else 0
                producto_existente.utilidad = float(utilidad_bs) if utilidad_bs else 0
                producto_existente.stock_minimo = int(stock_minimo) if stock_minimo else 0
                producto_existente.stock_maximo = int(stock_maximo) if stock_maximo else 0
                if estado:
                    producto_existente.activo = str(estado).strip().lower() == "activo"
                db.commit()
                resultados["actualizados"] += 1
            else:
                nuevo = Producto(
                    codigo=str(codigo).strip(),
                    categoria_id=categoria.id if categoria else None,
                    descripcion=str(descripcion),
                    marca=str(marca),
                    procedencia=str(procedencia) if procedencia else "",
                    precio=float(costo_bs) if costo_bs else 0,
                    utilidad=float(utilidad_bs) if utilidad_bs else 0,
                    stock_inicial=int(stock_inicial) if stock_inicial else 0,
                    stock_actual=int(stock_inicial) if stock_inicial else 0,
                    stock_minimo=int(stock_minimo) if stock_minimo else 0,
                    stock_maximo=int(stock_maximo) if stock_maximo else 0,
                    activo=str(estado).strip().lower() == "activo" if estado else True,
                    usuario_id=usuario_id,
                )
                db.add(nuevo)
                db.commit()
                resultados["creados"] += 1
        except Exception as e:
            db.rollback()
            resultados["errores"].append(f"Fila {row_num}: {str(e)}")
    
    broadcast_multiple_sync(["productos", "dashboard"], {"type": "updated", "room": "productos"})
    return resultados

@router.post("/delete-batch")
def delete_productos_batch_endpoint(ids: List[int], db: Session = Depends(get_db)):
    if not ids:
        raise HTTPException(status_code=400, detail="No se proporcionaron IDs")
    result = crud_delete_batch(db, ids)
    broadcast_multiple_sync(["productos", "dashboard"], {"type": "deleted", "room": "productos"})
    message = f"{result['hard_deleted']} eliminado(s), {result['soft_deleted']} desactivado(s) por estar en uso"
    return {"message": message, "count": result["count"], "soft_deleted": result["soft_deleted"], "hard_deleted": result["hard_deleted"]}

@router.delete("/all")
def delete_all_productos_endpoint(db: Session = Depends(get_db)):
    result = crud_delete_all(db)
    broadcast_multiple_sync(["productos", "dashboard"], {"type": "deleted", "room": "productos"})
    message = f"{result['hard_deleted']} eliminado(s), {result['soft_deleted']} desactivado(s) por estar en uso"
    return {"message": message, "count": result["count"], "soft_deleted": result["soft_deleted"], "hard_deleted": result["hard_deleted"]}

@router.get("/codigo/{codigo}", response_model=ProductoResponse)
def read_producto_by_codigo(codigo: str, db: Session = Depends(get_db)):
    db_producto = get_producto_by_codigo(db, codigo)
    if not db_producto:
        raise HTTPException(status_code=404, detail="Producto not found")
    from app.crud.producto import _build_full_response
    return _build_full_response(db, db_producto)

@router.get("/{producto_id}", response_model=ProductoResponse)
def read_producto(producto_id: int, db: Session = Depends(get_db)):
    result = get_producto_full(db, producto_id)
    if not result:
        raise HTTPException(status_code=404, detail="Producto not found")
    return result

@router.post("/", response_model=ProductoResponse)
def create_producto_endpoint(producto: ProductoCreate, db: Session = Depends(get_db), current_user=Depends(get_current_user_full)):
    db_producto = get_producto_by_codigo(db, producto.codigo)
    if db_producto:
        raise HTTPException(status_code=400, detail="Codigo already registered")
    result = create_producto(db, producto, usuario_id=current_user.id)
    broadcast_multiple_sync(["productos", "dashboard"], {"type": "created", "room": "productos"})
    return result

@router.put("/{producto_id}", response_model=ProductoResponse)
def update_producto_endpoint(producto_id: int, producto: ProductoUpdate, db: Session = Depends(get_db)):
    db_producto = update_producto(db, producto_id, producto)
    if not db_producto:
        raise HTTPException(status_code=404, detail="Producto not found")
    broadcast_multiple_sync(["productos", "dashboard"], {"type": "updated", "room": "productos"})
    return db_producto

@router.delete("/{producto_id}")
def delete_producto_endpoint(producto_id: int, db: Session = Depends(get_db)):
    result = delete_producto(db, producto_id)
    if not result:
        raise HTTPException(status_code=404, detail="Producto not found")
    broadcast_multiple_sync(["productos", "dashboard"], {"type": "deleted", "room": "productos"})
    if result["soft_deleted"]:
        return {"message": "Producto en uso, no se eliminó (se desactivó)", "id": producto_id, "soft_deleted": True, "en_uso": True}
    return {"message": "Producto eliminado", "id": producto_id, "soft_deleted": False, "en_uso": False}

@router.patch("/{producto_id}/toggle-activo", response_model=ProductoResponse)
def toggle_producto_activo(producto_id: int, db: Session = Depends(get_db)):
    from app.crud.producto import get_producto as _get_producto
    db_producto = _get_producto(db, producto_id)
    if not db_producto:
        raise HTTPException(status_code=404, detail="Producto not found")
    db_producto.activo = not db_producto.activo
    db.commit()
    db.refresh(db_producto)
    from app.crud.producto import _build_full_response
    broadcast_multiple_sync(["productos", "dashboard"], {"type": "updated", "room": "productos"})
    return _build_full_response(db, db_producto)

@router.post("/{producto_id}/imagen")
async def upload_producto_imagen(producto_id: int, file: UploadFile = File(...), db: Session = Depends(get_db)):
    db_producto = get_producto(db, producto_id)
    if not db_producto:
        raise HTTPException(status_code=404, detail="Producto not found")

    ext = os.path.splitext(file.filename or "")[1].lower()
    if ext not in ALLOWED_EXTENSIONS:
        raise HTTPException(status_code=400, detail=f"Formato no permitido: {ext or 'desconocido'}. Permitidos: {', '.join(sorted(ALLOWED_EXTENSIONS))}")

    UPLOAD_DIR.mkdir(parents=True, exist_ok=True)
    image_path = UPLOAD_DIR / f"{producto_id}{ext}"

    contents = await file.read()
    try:
        img = Image.open(BytesIO(contents))
        img.verify()
    except Exception:
        raise HTTPException(status_code=400, detail="El archivo no es una imagen válida")

    image_path.write_bytes(contents)
    url = f"/uploads/productos/{producto_id}{ext}"
    db_producto.imagen = url
    db.commit()
    db.refresh(db_producto)
    broadcast_multiple_sync(["productos", "dashboard"], {"type": "updated", "room": "productos"})
    return {"imagen": url}

@router.delete("/{producto_id}/imagen")
def delete_producto_imagen(producto_id: int, db: Session = Depends(get_db)):
    from app.crud.producto import get_producto as _get_producto
    db_producto = _get_producto(db, producto_id)
    if not db_producto:
        raise HTTPException(status_code=404, detail="Producto not found")
    if db_producto.imagen:
        try:
            old_path = Path(str(BASE_DIR / db_producto.imagen.lstrip("/")))
            if old_path.exists():
                old_path.unlink()
        except Exception:
            pass
        db_producto.imagen = None
        db.commit()
        db.refresh(db_producto)
        broadcast_multiple_sync(["productos", "dashboard"], {"type": "updated", "room": "productos"})
    return {"message": "Imagen eliminada", "imagen": None}
