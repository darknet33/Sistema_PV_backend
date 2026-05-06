import openpyxl
from openpyxl.styles import Font, PatternFill
from io import BytesIO
from typing import List, Dict

def export_to_excel(data: List[Dict], headers: List[str], sheet_name: str = "Sheet1"):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet_name
    
    header_font = Font(bold=True)
    header_fill = PatternFill(start_color="CCCCCC", end_color="CCCCCC", fill_type="solid")
    
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
    
    for row_idx, row_data in enumerate(data, 2):
        for col_idx, header in enumerate(headers, 1):
            ws.cell(row=row_idx, column=col_idx, value=row_data.get(header, ""))
    
    excel_file = BytesIO()
    wb.save(excel_file)
    excel_file.seek(0)
    return excel_file

def import_from_excel(file_path: str) -> List[Dict]:
    wb = openpyxl.load_workbook(file_path)
    ws = wb.active
    
    headers = [cell.value for cell in ws[1]]
    data = []
    
    for row in ws.iter_rows(min_row=2, values_only=True):
        if any(row):
            data.append(dict(zip(headers, row)))
    
    return data
